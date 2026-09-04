import os
import json
import time
import tempfile
from pathlib import Path

import streamlit as st
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document


# ============================================================
# CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)

# Primary model
PRIMARY_MODEL = "gemini-2.5-flash"

# Used automatically if the primary model is temporarily busy
FALLBACK_MODEL = "gemini-2.5-flash-lite"

# Maximum upload size
MAX_FILE_SIZE_MB = 10

# Number of attempts for each model
MAX_RETRIES_PER_MODEL = 3


# ============================================================
# API KEY
# ============================================================

def get_api_key():
    """
    Get Gemini API key from:
    1. Streamlit Secrets
    2. Environment variable
    """

    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        key = ""

    if key:
        return key.strip()

    return os.getenv("GEMINI_API_KEY", "").strip()


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file_bytes):
    """
    Extract text from a text-based PDF.
    """

    reader = PdfReader(file_bytes)

    pages = []

    for page in reader.pages:
        pages.append(page.extract_text() or "")

    return "\n".join(pages).strip()


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_docx_text(file_bytes):
    """
    Extract text from DOCX paragraphs and tables.
    """

    with tempfile.NamedTemporaryFile(
        suffix=".docx",
        delete=False
    ) as temp_file:

        temp_file.write(file_bytes)
        temp_path = temp_file.name

    try:

        document = Document(temp_path)

        parts = []

        # Paragraphs
        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # Tables
        for table in document.tables:

            for row in table.rows:

                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                row_text = " | ".join(
                    cell for cell in cells if cell
                )

                if row_text:
                    parts.append(row_text)

        return "\n".join(parts).strip()

    finally:

        Path(temp_path).unlink(
            missing_ok=True
        )


# ============================================================
# RESUME TEXT EXTRACTION
# ============================================================

def extract_resume_text(uploaded_file):
    """
    Extract text from PDF, DOCX, or TXT.
    """

    file_bytes = uploaded_file.getvalue()

    extension = Path(
        uploaded_file.name
    ).suffix.lower()

    if extension == ".pdf":

        return extract_pdf_text(
            file_bytes
        )

    elif extension == ".docx":

        return extract_docx_text(
            file_bytes
        )

    elif extension == ".txt":

        return file_bytes.decode(
            "utf-8",
            errors="ignore"
        ).strip()

    else:

        raise ValueError(
            "Unsupported file type."
        )


# ============================================================
# CHECK TEMPORARY GEMINI ERRORS
# ============================================================

def is_temporary_gemini_error(error):
    """
    Detect temporary Gemini service errors such as:
    503 UNAVAILABLE
    429 RESOURCE EXHAUSTED
    high demand
    overloaded
    """

    error_text = str(error).lower()

    temporary_errors = [

        "503",

        "unavailable",

        "high demand",

        "overloaded",

        "429",

        "resource exhausted",

        "deadline exceeded",

        "temporarily unavailable",

    ]

    return any(
        error_type in error_text
        for error_type in temporary_errors
    )


# ============================================================
# GEMINI RESUME ANALYSIS
# ============================================================

def analyze_resume(
    resume_text,
    job_description
):
    """
    Analyze the resume using Gemini.

    The function:
    1. Tries Gemini Flash
    2. Retries temporary failures
    3. Falls back to Gemini Flash Lite
    """

    api_key = get_api_key()

    if not api_key:

        raise RuntimeError(
            "Gemini API key is missing.\n\n"
            "Add GEMINI_API_KEY to Streamlit Secrets."
        )

    client = genai.Client(
        api_key=api_key
    )

    # --------------------------------------------------------
    # JSON RESPONSE SCHEMA
    # --------------------------------------------------------

    response_schema = {

        "type": "object",

        "properties": {

            "ats_score": {
                "type": "integer",
                "description":
                    "Estimated ATS compatibility score from 0 to 100."
            },

            "summary": {
                "type": "string"
            },

            "strengths": {

                "type": "array",

                "items": {
                    "type": "string"
                }
            },

            "missing_keywords": {

                "type": "array",

                "items": {
                    "type": "string"
                }
            },

            "formatting_issues": {

                "type": "array",

                "items": {
                    "type": "string"
                }
            },

            "improvements": {

                "type": "array",

                "items": {
                    "type": "string"
                }
            },

            "action_plan": {

                "type": "array",

                "items": {
                    "type": "string"
                }
            }
        },

        "required": [

            "ats_score",

            "summary",

            "strengths",

            "missing_keywords",

            "formatting_issues",

            "improvements",

            "action_plan"

        ]
    }

    # --------------------------------------------------------
    # LIMIT INPUT SIZE
    # --------------------------------------------------------

    resume_text = resume_text[:50000]

    job_text = (
        job_description.strip()
        if job_description.strip()
        else "No job description was provided."
    )

    job_text = job_text[:30000]

    # --------------------------------------------------------
    # PROMPT
    # --------------------------------------------------------

    prompt = f"""

You are an expert Resume Reviewer and ATS Optimization Specialist.

Analyze the resume and estimate how well it would perform in an
Applicant Tracking System (ATS).

IMPORTANT:

- This is an estimated ATS compatibility score.
- It is NOT an official score from a specific ATS company.
- Score from 0 to 100.
- Do NOT invent experience.
- Do NOT invent skills.
- Do NOT invent education.
- Do NOT invent certifications.
- Do NOT invent employers.
- Do NOT invent dates.
- Do NOT invent achievements.
- Do NOT invent keywords as if they exist in the resume.

If a job description is provided, compare the resume against it.

If no job description is provided, evaluate general ATS readiness.

SCORING:

0-25   = Very weak
26-50  = Needs major improvement
51-70  = Fair
71-85  = Strong
86-100 = Excellent

Evaluate:

1. Job-description keyword alignment
2. Relevant technical and soft skills
3. Standard resume sections
4. Measurable achievements
5. Clear job titles
6. Employment dates
7. ATS-safe formatting
8. Resume readability
9. Action verbs
10. Relevance of work experience
11. Professional summary
12. Education section
13. Skills section

Give specific and practical recommendations.

JOB DESCRIPTION:

{job_text}


RESUME:

{resume_text}

"""

    # --------------------------------------------------------
    # GEMINI CONFIGURATION
    # --------------------------------------------------------

    generation_config = types.GenerateContentConfig(

        temperature=0.2,

        response_mime_type="application/json",

        response_schema=response_schema,

    )

    # --------------------------------------------------------
    # TRY PRIMARY + FALLBACK MODELS
    # --------------------------------------------------------

    models_to_try = [

        PRIMARY_MODEL,

        FALLBACK_MODEL

    ]

    last_error = None

    for model_name in models_to_try:

        for attempt in range(
            1,
            MAX_RETRIES_PER_MODEL + 1
        ):

            try:

                response = client.models.generate_content(

                    model=model_name,

                    contents=prompt,

                    config=generation_config,

                )

                # Check empty response

                if not response.text:

                    raise RuntimeError(
                        f"{model_name} returned an empty response."
                    )

                # Parse JSON

                result = json.loads(
                    response.text
                )

                # Make sure ATS score is valid

                score = int(
                    result.get(
                        "ats_score",
                        0
                    )
                )

                result["ats_score"] = max(
                    0,
                    min(
                        100,
                        score
                    )
                )

                # Store model internally

                result["_model_used"] = model_name

                return result

            except Exception as error:

                last_error = error

                # Only retry temporary service errors

                if not is_temporary_gemini_error(
                    error
                ):

                    raise

                # Retry

                if attempt < MAX_RETRIES_PER_MODEL:

                    # Wait:
                    # attempt 1 → 2 seconds
                    # attempt 2 → 4 seconds

                    wait_time = 2 ** attempt

                    time.sleep(
                        wait_time
                    )

    # Both models failed

    raise RuntimeError(

        "Gemini is temporarily unavailable after "
        "multiple retries and fallback attempts.\n\n"
        "Please wait a few minutes and try again.\n\n"
        f"Last error: {last_error}"

    )


# ============================================================
# DISPLAY LIST
# ============================================================

def display_list(
    items,
    empty_message
):

    if items:

        for item in items:

            st.write(
                f"- {item}"
            )

    else:

        st.success(
            empty_message
        )


# ============================================================
# MAIN APP
# ============================================================

def main():

    # --------------------------------------------------------
    # HEADER
    # --------------------------------------------------------

    st.title(
        "📄 Resume ATS Analyzer"
    )

    st.write(
        "Upload your resume to get an estimated ATS score, "
        "keyword gaps, formatting issues, and actionable improvements."
    )

    # --------------------------------------------------------
    # SIDEBAR
    # --------------------------------------------------------

    with st.sidebar:

        st.header(
            "⚙️ Settings"
        )

        st.write(
            f"**Primary model:** `{PRIMARY_MODEL}`"
        )

        st.write(
            f"**Fallback model:** `{FALLBACK_MODEL}`"
        )

        st.write(
            "**Supported:** PDF, DOCX, TXT"
        )

        st.write(
            f"**Maximum file size:** "
            f"{MAX_FILE_SIZE_MB} MB"
        )

        st.divider()

        st.info(
            "For the most accurate ATS analysis, "
            "paste the target job description."
        )

        st.warning(
            "Never put your Gemini API key directly "
            "inside app.py or GitHub."
        )

    # --------------------------------------------------------
    # RESUME UPLOAD
    # --------------------------------------------------------

    uploaded_file = st.file_uploader(

        "📤 Upload your resume",

        type=[
            "pdf",
            "docx",
            "txt"
        ],

        help=(
            "Upload a text-based PDF, DOCX, "
            "or TXT resume."
        )

    )

    # --------------------------------------------------------
    # JOB DESCRIPTION
    # --------------------------------------------------------

    job_description = st.text_area(

        "💼 Target Job Description",

        height=220,

        placeholder=(
            "Paste the job description here. "
            "The AI will compare your resume against "
            "the job requirements."
        )

    )

    # --------------------------------------------------------
    # NO FILE
    # --------------------------------------------------------

    if not uploaded_file:

        st.info(
            "📌 Upload your resume above to start."
        )

        return

    # --------------------------------------------------------
    # FILE SIZE
    # --------------------------------------------------------

    file_size_mb = (
        uploaded_file.size
        /
        (1024 * 1024)
    )

    st.caption(

        f"Selected file: "
        f"**{uploaded_file.name}** · "
        f"{file_size_mb:.2f} MB"

    )

    if file_size_mb > MAX_FILE_SIZE_MB:

        st.error(

            f"File is too large. "
            f"Please upload a file smaller than "
            f"{MAX_FILE_SIZE_MB} MB."

        )

        return

    # --------------------------------------------------------
    # ANALYZE BUTTON
    # --------------------------------------------------------

    if st.button(

        "🔍 Analyze Resume",

        type="primary",

        use_container_width=True

    ):

        try:

            # Extract text

            with st.spinner(
                "📖 Reading your resume..."
            ):

                resume_text = extract_resume_text(
                    uploaded_file
                )

            # Empty resume

            if not resume_text:

                st.error(
                    "No readable text was found in the resume."
                )

                st.info(
                    "If you uploaded a scanned PDF, "
                    "try a text-based PDF or DOCX file."
                )

                return

            # Gemini analysis

            with st.spinner(

                "🤖 Gemini is analyzing your resume..."

            ):

                result = analyze_resume(

                    resume_text=resume_text,

                    job_description=job_description

                )

            # Save result

            st.session_state[
                "analysis"
            ] = result

        except Exception as error:

            error_text = str(
                error
            )

            # ------------------------------------------------
            # 503 / 429 ERROR
            # ------------------------------------------------

            if is_temporary_gemini_error(
                error
            ):

                st.error(
                    "⚠️ Gemini is temporarily busy."
                )

                st.info(

                    "The app already retried the request "
                    "and tried the fallback model. "
                    "Please wait a few minutes and click "
                    "**Analyze Resume** again."

                )

            # ------------------------------------------------
            # OTHER ERROR
            # ------------------------------------------------

            else:

                st.error(
                    f"❌ Analysis failed: {error_text}"
                )

                st.info(

                    "Please check your Gemini API key, "
                    "internet connection, file format, "
                    "and Streamlit logs."

                )

    # --------------------------------------------------------
    # GET SAVED RESULT
    # --------------------------------------------------------

    result = st.session_state.get(
        "analysis"
    )

    if not result:

        return

    # --------------------------------------------------------
    # RESULTS
    # --------------------------------------------------------

    st.divider()

    score_column, summary_column = st.columns(
        [1, 2]
    )

    # --------------------------------------------------------
    # ATS SCORE
    # --------------------------------------------------------

    with score_column:

        score = result[
            "ats_score"
        ]

        st.metric(

            "Estimated ATS Score",

            f"{score}/100"

        )

        st.progress(
            score / 100
        )

        if score >= 86:

            st.success(
                "🌟 Excellent ATS readiness"
            )

        elif score >= 71:

            st.success(
                "✅ Strong ATS readiness"
            )

        elif score >= 51:

            st.warning(
                "⚠️ Fair ATS readiness"
            )

        else:

            st.error(
                "❌ Needs significant improvement"
            )

        model_used = result.get(
            "_model_used"
        )

        if model_used:

            st.caption(
                f"Analysis generated with "
                f"`{model_used}`"
            )

    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    with summary_column:

        st.subheader(
            "📋 Overall Assessment"
        )

        st.write(
            result[
                "summary"
            ]
        )

    # --------------------------------------------------------
    # STRENGTHS
    # --------------------------------------------------------

    st.subheader(
        "💪 Strengths"
    )

    display_list(

        result.get(
            "strengths",
            []
        ),

        "No major strengths were identified."

    )

    # --------------------------------------------------------
    # MISSING KEYWORDS
    # --------------------------------------------------------

    st.subheader(
        "🔑 Missing / Weak Keywords"
    )

    display_list(

        result.get(
            "missing_keywords",
            []
        ),

        "No major keyword gaps were identified."

    )

    # --------------------------------------------------------
    # FORMATTING ISSUES
    # --------------------------------------------------------

    st.subheader(
        "⚠️ Formatting Issues"
    )

    display_list(

        result.get(
            "formatting_issues",
            []
        ),

        "No major formatting issues were identified."

    )

    # --------------------------------------------------------
    # IMPROVEMENTS
    # --------------------------------------------------------

    st.subheader(
        "🛠️ Recommended Improvements"
    )

    improvements = result.get(
        "improvements",
        []
    )

    if improvements:

        for index, item in enumerate(
            improvements,
            1
        ):

            st.write(
                f"**{index}.** {item}"
            )

    else:

        st.success(
            "No additional improvements were identified."
        )

    # --------------------------------------------------------
    # ACTION PLAN
    # --------------------------------------------------------

    st.subheader(
        "🚀 Action Plan"
    )

    action_plan = result.get(
        "action_plan",
        []
    )

    if action_plan:

        for index, item in enumerate(
            action_plan,
            1
        ):

            st.write(
                f"**{index}.** {item}"
            )

    else:

        st.success(
            "No additional action items were identified."
        )

    # --------------------------------------------------------
    # DOWNLOAD JSON
    # --------------------------------------------------------

    download_result = {

        key: value

        for key, value in result.items()

        if not key.startswith("_")

    }

    analysis_json = json.dumps(

        download_result,

        indent=2,

        ensure_ascii=False

    )

    st.download_button(

        "⬇️ Download Analysis as JSON",

        data=analysis_json,

        file_name="resume_ats_analysis.json",

        mime="application/json"

    )


# ============================================================
# RUN APPLICATION
# ============================================================

if __name__ == "__main__":

    main()
```
